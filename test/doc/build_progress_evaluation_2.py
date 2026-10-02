"""Build the PE2 deck from recorded project results; never retrain the model."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

ROOT = Path(__file__).resolve().parent
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BG, INK, GREEN, MUTED = 'F4F7F5', '153C36', '19806B', '536963'
outline = []

def text(slide, x, y, w, h, value, size=22, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Aptos'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(14)
    return shape

def slide(title, owner, body, notes, source='Model/new.ipynb · TEAM_VIVA_GUIDE.md'):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    text(s, .65, .35, 12, .4, 'STAYWISE / PROGRESS EVALUATION 2 / IT3051', 12, GREEN, True)
    text(s, .65, 1, 12, 1.1, title, 34, INK, True)
    if body:
        text(s, .8, 2.35, 11.8, 4.2, body, 23)
    text(s, .65, 6.83, 11.7, .35, f'{owner}  |  {source}', 10, MUTED)
    text(s, 12.1, 6.83, .6, .35, str(len(prs.slides)).zfill(2), 11, GREEN)
    s.notes_slide.notes_text_frame.text = notes
    outline.append((title, owner, body, notes))
    return s

def table(s, headers, rows, y=2.35, widths=None):
    shape = s.shapes.add_table(len(rows)+1, len(headers), Inches(.8), Inches(y), Inches(11.75), Inches(.56*(len(rows)+1)))
    t = shape.table
    if widths:
        for c,w in zip(t.columns,widths): c.width= Inches(w)
    for i,row in enumerate([headers]+rows):
        for j,value in enumerate(row):
            c=t.cell(i,j)
            c.text=str(value)
            c.fill.solid()
            c.fill.fore_color.rgb=RGBColor.from_string(INK if i==0 else ('E2EEE8' if i%2 else 'FFFFFF'))
            for p in c.text_frame.paragraphs:
                p.font.name='Aptos'; p.font.size=Pt(18); p.font.bold=(i==0)
                p.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else INK)

slide('Hotel booking cancellation prediction', 'Team introduction',
      'Model development, optimization and final selection\nSupervised binary classification: 0 = kept; 1 = canceled\nManoharan Subasthican · IT23810778\nKuganesan Kageepan · IT23717572\nGowsika Surendran · IT23794566\nKokulan Seyon · IT23752672',
      'Introduce the team. PE2 is an individual evaluation worth 30%, focused on official Stages 6 and 7. This deck supports the viva; every member must understand the complete workflow.', 'Official assignment · Stage 8')
slide('The modelling question', 'Subasthican',
      'Can booking details distinguish cancellations from kept bookings?\nUsers: reservation staff, hotel managers and operational planners.\nCompare four algorithms under the same validation conditions.\nTune suitable models and justify the final choice using evidence.',
      'Describe the cancellation risk estimate as decision support. Do not claim measured financial savings or guaranteed outcomes. Explain the difference between the modelling evaluation and the final business demonstration.')
slide('Data carried forward from Evaluation 1', 'Gowsika',
      '119,390 raw bookings → 87,396 after removing 31,994 duplicates.\nCleaned cancellation rate: 27.49% (raw rate: 37.04%).\n33 predictors: 21 numeric + 12 categorical.\nOutcome fields removed: reservation_status and reservation_status_date.\nSource period: July 2015–August 2017; two hotels.',
      'Explain why removing duplicates changes the class balance. Children use zero, country Unknown, and missing agent/company zero; identifiers become categorical strings for modelling. ADR outliers were investigated but not clipped. Source: Mostipak Kaggle distribution and Antonio et al. (2019).')
slide('Validation protects the test set', 'Seyon',
      'Stratified 80/20 split, random seed 42.\nTraining: 69,916 rows | Test: 17,480 rows.\nThree shuffled stratified folds within the training partition.\nEach fold fits its own preprocessing and classifier pipeline.\nSelect by training CV ROC-AUC → refit → evaluate on test.',
      'Test labels: 12,675 kept and 4,805 canceled. Cross-validation alternates training and validation folds; the separate test set is not a tuning input. Numeric median imputation and StandardScaler; categorical imputation and one-hot encoding with infrequent-category grouping at min_frequency=50. Stratification preserves class proportions. Random splitting does not prove future-period performance.')
s=slide('Four algorithms, four useful comparisons', 'All members', '',
      'Each member explains the algorithm allocated in the team guide. Ensembles can capture interactions; their empirical advantage is supported by validation, while precise causal reasons for a score difference are not established.')
table(s,['Algorithm','Why compare it?','Baseline settings'],[
    ['Logistic Regression','Linear probability baseline','max_iter = 1000'],
    ['Decision Tree','Nonlinear decision rules','depth = 20; leaf ≥ 5'],
    ['Random Forest','Average many fitted trees','120 trees; leaf ≥ 2'],
    ['Extra Trees','More randomized tree splits','120 trees; leaf ≥ 2']],widths=[2.7,4.5,4.55])
slide('Metrics answer different questions', 'Kageepan',
      'ROC-AUC: ranking across thresholds; primary selection metric.\nPrecision: how many cancellation alerts are correct?\nRecall: how many actual cancellations are found?\nF1: balance between cancellation precision and recall.\nAccuracy / balanced accuracy / average precision: supporting views.\nAlways predicting kept gives about 72.51% accuracy.',
      'Positive class is cancellation. Precision = TP/(TP+FP); recall = TP/(TP+FN); F1 is their harmonic mean. Balanced accuracy averages class recalls. Average precision summarizes the precision-recall curve. ROC-AUC is not accuracy, and a high AUC does not establish calibrated probabilities.')
s=slide('Baseline comparison: same folds and pipeline', 'Subasthican', '',
      'All values are three-fold training CV means, rounded to four decimals. Random Forest ranks first by ROC-AUC; Decision Tree has greater recall than baseline Forest but lower precision and AUC. This illustrates why the selection metric must be stated before comparing models.')
table(s,['Model','ROC-AUC','Accuracy','Precision','Recall','F1'],[
    ['Random Forest','.9097','.8474','.7887','.6075','.6863'],
    ['Extra Trees','.9058','.8435','.7802','.5994','.6779'],
    ['Logistic Regression','.8640','.8062','.6893','.5372','.6038'],
    ['Decision Tree','.8476','.8166','.6744','.6439','.6587']],widths=[3,1.75,1.75,1.75,1.75,1.75])
text(s,.85,5.65,11.4,.65,'Random Forest leads on ranking; no single model leads on every metric.',21,GREEN,True)
s=slide('Validation ranking and overfitting signals', 'Seyon', '',
      'The horizontal axis starts at zero. Baseline train minus validation AUC gaps are LR .0027, DT .1144, RF .0762, ET .0811. A larger gap suggests overfitting, but does not alone decide the winner. Fold AUC standard deviations are .0026, .0025, .0015 and .0017 respectively; these are not confidence intervals.')
data=CategoryChartData(); data.categories=['Logistic Regression','Decision Tree','Random Forest','Extra Trees']; data.add_series('Mean CV ROC-AUC',[.8640,.8476,.9097,.9058])
chart=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(.8), Inches(2.2), Inches(7.4), Inches(3.9),data).chart
chart.has_legend=False; chart.value_axis.minimum_scale=0; chart.value_axis.maximum_scale=1
chart.plots[0].has_data_labels=True; chart.plots[0].data_labels.number_format='0.0000'; chart.plots[0].data_labels.position=XL_LABEL_POSITION.OUTSIDE_END
chart.series[0].format.fill.solid(); chart.series[0].format.fill.fore_color.rgb=RGBColor.from_string(GREEN)
text(s,8.55,2.45,4,3.7,'Train–validation AUC gap\nDecision Tree: 0.1144\nRandom Forest: 0.0762\nExtra Trees: 0.0811\nLogistic: 0.0027',20)
slide('Random Forest: constrained randomized search', 'Gowsika',
      'RandomizedSearchCV: 4 sampled candidates × 3 folds.\nSearch: trees {100,160}; depth {None,20}; leaf {2,5,10}.\nFeature sampling: {sqrt,0.7}; scoring = ROC-AUC.\nBest: 100 trees, depth 20, leaf size 2, feature fraction 0.7.\nCV ROC-AUC: 0.9097 → 0.9125 (+0.0028).',
      'The full parameter space has 24 combinations, but only four were sampled with seed 42. Do not call this exhaustive search. Depth and minimum leaf size control complexity. n_jobs=-1 inside Forest and search n_jobs=1 avoid nested parallelism. The limited budget is a limitation.')
slide('Logistic Regression: small exhaustive grid', 'Subasthican',
      'GridSearchCV: 6 candidates × 3 folds.\nC ∈ {0.3, 1.0, 3.0}; class_weight ∈ {None, balanced}.\nSmaller C means stronger regularization.\nBest: C = 3.0; balanced class weights.\nCV ROC-AUC: 0.8640 → 0.8644 (+0.0004).',
      'Balanced class weights alter the training loss inversely to class frequency; they are not oversampling or SMOTE. Same folds and scoring as Forest. A small improvement is observed; no statistical significance claim is supported.')
s=slide('Final selection uses training evidence', 'Subasthican', '',
      'Rank all six candidates by training CV AUC. Clone the top candidate and refit on all training rows before testing. The choice follows the declared criterion rather than test-set comparison. The tuning estimates themselves may be optimistic because the same CV supports parameter selection.')
table(s,['Candidate','Training CV ROC-AUC'],[
    ['Tuned Random Forest','0.9125'],['Baseline Random Forest','0.9097'],['Baseline Extra Trees','0.9058'],
    ['Tuned Logistic Regression','0.8644'],['Baseline Logistic Regression','0.8640'],['Baseline Decision Tree','0.8476']],widths=[8,3.75])
slide('Feature ablation finds no useful gain', 'Gowsika',
      'Derived: total_nights; is_long_stay (>7 nights).\nDerived: is_weekend_stay (>0); is_high_lead_time (>60 days).\nWith engineering: CV ROC-AUC 0.9125.\nWithout engineering: CV ROC-AUC 0.9126.\nDifference ≈ −0.0001: no demonstrated improvement.\nThe supplied pipeline retains the four features.',
      'The ablation removes all four derived columns using fixed selected Forest settings and the same folds. It does not retune both representations. The tiny difference is not evidence of significant harm or benefit. Deployment preparation must reproduce the retained feature rules exactly.')
s=slide('Held-out results: 17,480 unseen bookings', 'Kageepan', '',
      'These are recorded full-data held-out metrics, not reduced smoke-test scores. Test evaluation follows selection. Threshold 0.5. Accuracy exceeds the all-kept baseline by about 12.57 percentage points, but recall and precision expose important errors. Saved artifact could not be independently loaded during deck preparation; results were cross-checked against notebook outputs and documented tables.')
table(s,['Metric','Score','Interpretation'],[
    ['ROC-AUC','0.9158','Ranking discrimination'],['Accuracy','85.08%','All correct predictions'],
    ['Precision','75.43%','Correct cancellation alerts'],['Recall','67.80%','Actual cancellations detected'],
    ['F1','71.42%','Precision–recall balance'],['Balanced accuracy','79.72%','Average class recall']],widths=[3.4,2.2,6.15])
s=slide('The errors matter for hotel operations', 'Kageepan', '',
      'Rows are actual outcomes and columns are predicted outcomes. TP=3258, FN=1547, FP=1061, TN=11614. Precision 3258/(3258+1061); recall 3258/4805. Missing cancellations affects planning; false alerts require unnecessary follow-up. Decisions remain with staff.')
table(s,['Actual outcome','Predicted kept','Predicted canceled'],[
    ['Kept','11,614 · true negatives','1,061 · false positives'],
    ['Canceled','1,547 · false negatives','3,258 · true positives']],widths=[3.3,4.2,4.25])
text(s,.85,4.65,11.5,1.7,'1,547 cancellations missed | 1,061 unnecessary alerts\nUse risk estimates to prioritize courteous follow-up and occupancy review.\nThreshold 0.5 is not optimized for business costs.',22)
slide('Limits and the next modelling experiments', 'Gowsika / Seyon',
      'Historical two-hotel data may differ from current local operations.\nRandom split: future-period generalization remains untested.\nAssigned rooms and booking changes need a defined prediction time.\nLimited tuning; probabilities are not calibrated.\nNext: temporal/external validation, calibration and cost-based thresholds.',
      'Explain prediction-time leakage beyond explicitly removed status fields. A field can be valid for late-stage assessment yet unavailable when the booking is first made. Fairness/subgroup performance and current-hotel impact remain unmeasured. These are proposed experiments, not completed results.')
slide('Individual preparation and notebook evidence', 'All members',
      'Subasthican: Logistic Regression, tuning, candidate comparison.\nKageepan: Decision Tree, metrics and held-out error interpretation.\nGowsika: Random Forest, tuning and feature ablation.\nSeyon: pipeline, split/CV, Extra Trees and artifact export.\nEach member: explain one decision, its evidence and its limitation.',
      'This is a study/speaking allocation, not proof of original authorship. Confirm actual personal contributions and identify exact notebook cells. Suggested core speaking allocation: approximately four minutes each, adjusted to the evaluator’s format. Everyone must be ready to explain the full workflow.')
slide('Evidence and questions', 'All members',
      'Model/new.ipynb: baselines, searches, selection, ablation and test outputs.\nTEAM_VIVA_GUIDE.md: verified result tables and member preparation.\nOfficial assignment: Stage 8, Progress Evaluation 2 requirements.\nDataset: Mostipak, Hotel Booking Demand, Kaggle.\nOriginal publication: Antonio, de Almeida & Nunes (2019).\nDOI: 10.1016/j.dib.2018.11.126',
      'Kaggle source recorded in the proposal: https://www.kaggle.com/datasets/jessemostipak/hotel-booking-demand. Original publication: https://doi.org/10.1016/j.dib.2018.11.126. This deck uses local project evidence, not a new training run. Backup viva questions: Why AUC? Why stratification? Why a pipeline? Does engineering improve the model? Why not choose on test accuracy?', 'Local project evidence')

out = ROOT / 'StayWise_Progress_Evaluation_2.pptx'
prs.save(out)
lines=['# Progress Evaluation 2 — speaker guide','', 'Editable deck: StayWise_Progress_Evaluation_2.pptx', '',
       'Suggested core delivery: about 16 minutes, with roughly four minutes per member. Adjust to the evaluator’s time limit. Shared slides can be brief introductions or viva references.', '',
       'The allocation below follows TEAM_VIVA_GUIDE.md and describes speaking/study responsibility. Confirm actual individual contributions before presenting.', '',
       '**Evidence limitation:** the current model file raises a zlib decompression error when loaded. The presentation reports existing full-data notebook/document results; it does not claim the artifact was freshly revalidated. Resolve the artifact before any live prediction demonstration.', '']
for i,(title,owner,body,notes) in enumerate(outline,1):
    lines += [f'## Slide {i}: {title}',f'Presenter: {owner}','',notes,'']
(ROOT/'PROGRESS_EVALUATION_2_SPEAKER_GUIDE.md').write_text('\n'.join(lines),encoding='utf-8')
check=Presentation(out)
assert len(check.slides)==17
assert all(s.has_notes_slide and s.notes_slide.notes_text_frame.text for s in check.slides)
for s in check.slides:
    for sh in s.shapes:
        assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=prs.slide_width and sh.top+sh.height<=prs.slide_height
print(f'Created and validated {len(check.slides)} slides with notes: {out}')
