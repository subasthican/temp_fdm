"""Generate editable report and presentation from the project's documented evidence."""
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches as DocInches, Pt as DocPt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

HERE = Path(__file__).resolve().parent
document = Document()
normal = document.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = DocPt(11)
table = None
for line in (HERE / 'STAGE_9_12_REPORT.md').read_text(encoding='utf-8').splitlines():
    if not line.strip():
        table = None
        continue
    if line.startswith('|'):
        if re.match(r'^\|[\s|:-]+$', line):
            continue
        cells = [v.strip() for v in line.strip('|').split('|')]
        if table is None:
            table = document.add_table(rows=0, cols=len(cells))
            table.style = 'Light Shading Accent 1'
        for cell, text in zip(table.add_row().cells, cells):
            cell.text = text
    elif line.startswith('# '):
        document.add_heading(line[2:], 0)
    elif line.startswith('## '):
        document.add_heading(line[3:], 1)
    elif line.startswith('- '):
        document.add_paragraph(line[2:], style='List Bullet')
    else:
        document.add_paragraph(line.replace('`', ''))
section = document.sections[0]
section.header.paragraphs[0].text = 'IT3051 | StayWise | Technical Report Draft'
section.footer.paragraphs[0].text = 'Review team contributions and institutional details before submission.'
if (HERE / 'screenshots' / 'desktop-result.png').exists():
    document.add_heading('Appendix: Application and browser verification', 1)
    document.add_paragraph('The real application passed desktop and mobile browser checks for predictions, assessment downloads, field errors, stale-result clearing and reset. See SYSTEM_TEST_RESULTS.md for reproducible evidence.')
    document.add_picture(str(HERE / 'screenshots' / 'desktop-result.png'), width=DocInches(6))
document.save(HERE / 'StayWise_Technical_Report.docx')

slides = [
    ('More predictable stays.', 'STAYWISE · HOTEL BOOKING CANCELLATION', ['A practical decision-support system for hotel reservations.', 'Group presentation · IT3051 · October 2026'], 'Introduce the group and state the problem in plain language. Confirm institution and presenter names before submission.'),
    ('A booking is not guaranteed occupancy.', 'THE BUSINESS PROBLEM', ['Cancellations change expected room demand.', 'Uncertainty complicates staffing and room planning.', 'Reservation teams need a way to prioritize follow-up.'], 'Explain a realistic example without inventing financial loss figures.'),
    ('Built for reservation teams.', 'WHO USES THE SOLUTION?', ['Staff enter the details currently known about a booking.', 'Managers receive a cancellation estimate and next action.', 'People remain responsible for customer decisions.'], 'The current model uses room assignments and booking changes, so discuss when those details are available.'),
    ('Historical bookings provide the evidence.', 'DATA WE ANALYSED', ['119,390 original records · City Hotel and Resort Hotel.', '87,396 records remain after exact duplicate removal.', '27.49% canceled in the cleaned data.', 'Arrivals span July 2015 to August 2017.'], 'Credit Mostipak’s Kaggle distribution and Antonio, de Almeida and Nunes (2019), DOI 10.1016/j.dib.2018.11.126. Raw cancellation rate is 37.04%; do not confuse the two populations.'),
    ('A consistent path from details to risk.', 'HOW THE SOLUTION WORKS', ['Enter booking details → validate the input.', 'Apply the same preparation as the trained model.', 'Return estimated risk, likely outcome and follow-up.', 'Download the assessment for demonstration records.'], 'Explain no retraining occurs during a request. Four features are computed: total nights, long stay, weekend stay and high lead time.'),
    ('We selected the strongest validation result.', 'MODEL CHOICE', ['Compared four approaches on the same training folds.', 'Tuned Random Forest: validation ranking score 0.9125.', 'Baseline Random Forest: 0.9097; Extra Trees: 0.9058.', 'An untouched test partition checks the final choice.'], 'ROC-AUC measures ranking; it is not percentage accuracy. Logistic Regression and Decision Tree were also compared. Engineering did not improve validation AUC.'),
    ('Useful evidence, with visible mistakes.', 'HELD-OUT RESULTS', ['85.1% overall accuracy on 17,480 test bookings.', 'Identified 3,258 of 4,805 actual cancellations.', 'Missed 1,547 cancellations.', 'Flagged 1,061 bookings that were actually kept.'], 'Recall is 67.8% and precision is 75.4%. Explain why staff judgement matters.'),
    ('See a booking assessment in action.', 'LIVE DEMONSTRATION', ['Load the example booking and assess its risk.', 'Explain probability, outcome and suggested follow-up.', 'Change an input and reassess; show a validation error.', 'Download the assessment as JSON.'], 'Open http://127.0.0.1:5000. Do not promise that changing one input must increase risk. Demonstrate zero guests and a required missing field.'),
    ('Turn risk into proportionate action.', 'BUSINESS RECOMMENDATION', ['Higher risk: prioritize a courteous confirmation.', 'Moderate risk: monitor changes and send reminders.', 'Lower risk: continue routine follow-up.', 'Pilot the workflow before changing business policies.'], 'Risk bands are interface conventions; predictions do not justify denying a booking. No measured financial gain is claimed.'),
    ('Reliable software is part of the solution.', 'SYSTEM VALIDATION', ['Seven automated service tests pass.', 'Twenty historical bookings match direct model predictions.', 'Invalid values and missing fields receive clear errors.', 'Frontend and backend run together on one local service.'], 'Show tests/test_app.py and SYSTEM_TEST_RESULTS.md. Model accuracy comes from the notebook, not the software tests.'),
    ('Validate locally before wider use.', 'LIMITS & NEXT STEPS', ['Historical data may differ from today’s hotel operations.', 'Scores are not calibrated guarantees.', 'New hotels and identifiers require validation.', 'Next: recent data, temporal tests and a monitored pilot.'], 'Discuss subgroup analysis, cost-based thresholds and production deployment as future work.'),
    ('Every member owns their decisions.', 'TEAM CONTRIBUTIONS & QUESTIONS', ['Manoharan Subasthican · IT23810778', 'Kuganesan Kageepan · IT23717572', 'Gowsika Surendran · IT23794566', 'Kokulan Seyon · IT23752672', 'Confirm each member’s actual work before submission.'], 'Use the team guide for study allocation, not evidence of code authorship. Disclose coding-assistant support according to institutional rules. Each member explains an actual contribution and decision.'),
]
presentation = Presentation()
presentation.slide_width = Inches(13.333)
presentation.slide_height = Inches(7.5)
for number, (title, eyebrow, bullets, notes) in enumerate(slides, 1):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('F5F7F3')
    def box(x, y, w, h, text, size, color, bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = 'Calibri'
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_after = Pt(20)
        return shape
    box(.8, .55, 11, .5, eyebrow, 13, '28654F', True)
    box(.8, 1.3, 11.6, 1.5, title, 38, '18352E', True)
    if number == 8 and (HERE / 'screenshots' / 'desktop-result.png').exists():
        box(.9, 3, 5.4, 3.6, '\n'.join('•  ' + bullet for bullet in bullets), 20, '53685C')
        slide.shapes.add_picture(str(HERE / 'screenshots' / 'desktop-result.png'), Inches(6.6), Inches(2.8), width=Inches(5.8))
    else:
        box(.9, 3, 11.4, 3.5, '\n'.join('•  ' + bullet for bullet in bullets), 23, '53685C')
    box(.8, 7, 10, .3, 'StayWise · Reservation intelligence · IT3051', 10, '72817B')
    box(12, 7, .6, .3, f'{number:02d}', 10, '72817B')
    slide.notes_slide.notes_text_frame.text = notes
presentation.save(HERE / 'StayWise_Presentation.pptx')
print('Generated editable report and 12-slide presentation.')
